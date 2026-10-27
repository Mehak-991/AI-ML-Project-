"""Loading CSV and Excel files into DataFrames with defensive limits."""

import csv
import warnings
import zipfile
from collections import Counter
from pathlib import Path, PurePath

import pandas as pd

from app.core.exceptions import DataValidationError, InvalidFileError, UnsupportedFileTypeError

# Extension -> format. Legacy .xls is deliberately unsupported (needs xlrd, old format).
SUPPORTED_EXTENSIONS: dict[str, str] = {".csv": "csv", ".xlsx": "excel"}

# Reject .xlsx archives that inflate far beyond their on-disk size (zip bombs).
MAX_XLSX_COMPRESSION_RATIO = 100


def detect_file_format(filename: str) -> str:
    suffix = PurePath(filename).suffix.lower()
    file_format = SUPPORTED_EXTENSIONS.get(suffix)
    if file_format is None:
        allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{suffix or '(none)'}'. Allowed: {allowed}."
        )
    return file_format


def _check_xlsx_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            uncompressed = sum(info.file_size for info in archive.infolist())
            names = set(archive.namelist())
    except zipfile.BadZipFile:
        raise InvalidFileError("The file is not a valid .xlsx workbook.") from None
    if "[Content_Types].xml" not in names:
        raise InvalidFileError("The file is not a valid .xlsx workbook.")
    if uncompressed > max(path.stat().st_size, 1) * MAX_XLSX_COMPRESSION_RATIO:
        raise InvalidFileError("The workbook's uncompressed size is suspiciously large.")


def _raw_header(path: Path, file_format: str) -> list[str]:
    """Header cells exactly as written (pandas silently renames duplicates to 'a.1')."""
    if file_format == "csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            return next(csv.reader(handle), [])
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        first_row = next(workbook.worksheets[0].iter_rows(max_row=1, values_only=True), ())
        return ["" if cell is None else str(cell) for cell in first_row]
    finally:
        workbook.close()


def _check_header(header: list[str]) -> None:
    names = [name.strip() for name in header]
    if not names:
        raise InvalidFileError("The file is empty or has no header row.")
    duplicates = sorted(name for name, count in Counter(names).items() if count > 1 and name)
    if duplicates:
        raise DataValidationError(
            "Duplicate column names in header.", details={"duplicate_columns": duplicates}
        )


def read_tabular_file(
    path: Path, file_format: str, *, max_rows: int, max_columns: int
) -> pd.DataFrame:
    """Read a CSV/XLSX file, enforcing row/column limits and a clean header."""
    try:
        if file_format == "excel":
            _check_xlsx_archive(path)
        _check_header(_raw_header(path, file_format))
        if file_format == "csv":
            # index_col=False stops pandas silently turning surplus fields into an index
            # (which shifts data into the wrong columns); its warning is made an error.
            with warnings.catch_warnings():
                warnings.simplefilter("error", pd.errors.ParserWarning)
                frame = pd.read_csv(
                    path,
                    nrows=max_rows + 1,
                    encoding="utf-8-sig",
                    low_memory=False,
                    index_col=False,
                )
        elif file_format == "excel":
            frame = pd.read_excel(path, sheet_name=0, nrows=max_rows + 1, engine="openpyxl")
        else:
            raise UnsupportedFileTypeError(f"Unsupported format '{file_format}'.")
    except (InvalidFileError, DataValidationError):
        raise
    except pd.errors.EmptyDataError:
        raise InvalidFileError("The file is empty.") from None
    except UnicodeDecodeError:
        raise InvalidFileError("CSV files must be UTF-8 encoded.") from None
    except pd.errors.ParserWarning:
        raise InvalidFileError("Some rows have more fields than the header.") from None
    except Exception as exc:  # untrusted input: any parser failure (e.g. XML ParseError from a
        # malformed or entity-bomb .xlsx) means an invalid file, not a server error.
        raise InvalidFileError(f"The file could not be parsed as {file_format.upper()}.") from exc

    frame.columns = [str(column).strip() for column in frame.columns]
    if len(frame) > max_rows:
        raise DataValidationError(f"The dataset exceeds the limit of {max_rows} rows.")
    if frame.shape[1] > max_columns:
        raise DataValidationError(f"The dataset exceeds the limit of {max_columns} columns.")
    return frame
